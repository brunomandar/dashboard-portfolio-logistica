# -*- coding: utf-8 -*-
"""
Gera os JSONs do Dashboard do Portfólio de Logística a partir dos Excel dos projetos.

Uso: duplo clique em atualizar_dashboard.bat (ou: python gerar_jsons.py)

O script:
  1. Lê todos os arquivos "LOG-xxx - Nome do projeto.xlsx" da pasta do SharePoint (sincronizada pelo OneDrive),
     incluindo extensões de projetos, como "LOG-010.1 - Nome.xlsx";
  2. Gera um json/LOG-xxx.json por projeto (ID sempre com 3 dígitos; extensões viram LOG-xxx.y.json);
  3. Recria projects.index.json, projects.index.js, portfolio.json, projects.data.js e MANIFESTO.txt;
  4. Remove da pasta json os LOG-xx.json antigos que não correspondem mais a nenhum Excel.
"""
import datetime as dt
import glob
import json
import os
import re
import sys
import unicodedata
import urllib.parse
import warnings

warnings.filterwarnings("ignore")

try:
    import openpyxl
except ImportError:
    print("ERRO: biblioteca openpyxl não instalada. Rode: python -m pip install openpyxl")
    sys.exit(1)

# ---------------------------------------------------------------------------
# CONFIGURAÇÃO
# ---------------------------------------------------------------------------
# Caminho da pasta 0_Portfolio_Logística sincronizada no computador.
# Deixe vazio ("") para o script procurar sozinho dentro da sua pasta de usuário.
PASTA_EXCEL = r""

NOME_PASTA_EXCEL = "0_Portfolio_Logística"
URL_PASTA_SHAREPOINT = (
    "https://corpclarobr.sharepoint.com/sites/USER-Transporte-TorrePerformanceeProjetos/"
    "Documentos Compartilhados/Sustentação Transporte/"
    "_Suporte operacional e melhoria contínua/_Portfolio/0_Portfolio_Logística/"
)
IGNORAR_IDS = {"LOG-000"}  # modelos

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PASTA_JSON = os.path.join(BASE_DIR, "json")


# ---------------------------------------------------------------------------
# Utilitários
# ---------------------------------------------------------------------------
def sem_acento(texto):
    return "".join(c for c in unicodedata.normalize("NFD", str(texto)) if unicodedata.category(c) != "Mn")


def chave(texto):
    return re.sub(r"[^a-z0-9]", "", sem_acento(texto).lower())


def limpar(valor):
    if valor is None:
        return ""
    if isinstance(valor, str):
        texto = valor.replace("\u200b", "").replace("\\n", "\n").strip()
        return "" if texto in ("-", "#VALUE!", "#N/A", "#REF!", "#DIV/0!") else texto
    return valor


def data_iso(valor):
    valor = limpar(valor)
    if isinstance(valor, dt.datetime):
        return valor.strftime("%Y-%m-%d")
    if isinstance(valor, dt.date):
        return valor.isoformat()
    if isinstance(valor, str):
        m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{2,4})$", valor)
        if m:
            d, mth, y = (int(x) for x in m.groups())
            y = y + 2000 if y < 100 else y
            return f"{y:04d}-{mth:02d}-{d:02d}"
        if re.match(r"^\d{4}-\d{2}-\d{2}", valor):
            return valor[:10]
    return ""


def numero(valor):
    valor = limpar(valor)
    if isinstance(valor, (int, float)):
        return float(valor)
    try:
        return float(str(valor).replace(",", "."))
    except (TypeError, ValueError):
        return 0.0


def inteiro_ou_texto(valor):
    valor = limpar(valor)
    if isinstance(valor, float) and valor.is_integer():
        return int(valor)
    return valor


def normalizar_id(texto):
    """LOG-15 -> LOG-015 | LOG-010.1 -> LOG-010.1 (extensão do projeto principal)."""
    m = re.search(r"LOG\s*-?\s*(\d+)(?:\.(\d+))?", str(texto), re.I)
    if not m:
        return ""
    principal = f"LOG-{int(m.group(1)):03d}"
    return f"{principal}.{int(m.group(2))}" if m.group(2) else principal


def ordem_id(project_id):
    """Ordena LOG-010, LOG-010.1, LOG-010.2, LOG-011..."""
    m = re.search(r"(\d+)(?:\.(\d+))?", project_id)
    return (int(m.group(1)), int(m.group(2) or 0)) if m else (0, 0)


def texto_celula(valor):
    valor = limpar(valor)
    if isinstance(valor, (dt.date, dt.datetime)):
        return data_iso(valor)
    return str(valor) if valor != "" else ""


# ---------------------------------------------------------------------------
# Leitura das abas
# ---------------------------------------------------------------------------
def mapa_rotulos(ws, linhas=range(1, 40), colunas=range(1, 9)):
    """Encontra células de rótulo (ex.: 'Projeto:') e devolve a posição de cada uma."""
    rotulos = {}
    for r in linhas:
        for c in colunas:
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip().endswith(":"):
                rotulos.setdefault(chave(v), (r, c))
    return rotulos


def valor_ao_lado(ws, rotulos, *nomes):
    for nome in nomes:
        pos = rotulos.get(chave(nome))
        if pos:
            return limpar(ws.cell(pos[0], pos[1] + 1).value)
    return ""


def valor_abaixo(ws, rotulos, *nomes):
    for nome in nomes:
        pos = rotulos.get(chave(nome))
        if pos:
            return limpar(ws.cell(pos[0] + 1, pos[1]).value)
    return ""


def ler_one_page(ws):
    rot = mapa_rotulos(ws)
    dados = {
        "projeto": valor_ao_lado(ws, rot, "Projeto:"),
        "area": valor_ao_lado(ws, rot, "Área:"),
        "pontoFocal": valor_ao_lado(ws, rot, "Ponto Focal:"),
        "inicio": data_iso(valor_ao_lado(ws, rot, "Início:")),
        "termino": data_iso(valor_ao_lado(ws, rot, "Término:")),
        "inicioCaptura": data_iso(valor_ao_lado(ws, rot, "Início Capt.:", "Início Captura:")),
        "iniciativa": valor_ao_lado(ws, rot, "Iniciativa:"),
        "pmo": valor_ao_lado(ws, rot, "Focal Acompanhamento:", "PMO:"),
        "risco": valor_ao_lado(ws, rot, "Risco:"),
        "status": valor_ao_lado(ws, rot, "Status:"),
        "prioridade": valor_ao_lado(ws, rot, "Prioridade:"),
        "ganhoProjeto": valor_ao_lado(ws, rot, "Ganho Projeto:"),
        "escopo": valor_abaixo(ws, rot, "Escopo:"),
        "entregaveisTexto": valor_abaixo(ws, rot, "Entregáveis do Projeto:"),
        "statusProjeto": valor_abaixo(ws, rot, "Status do Projeto:"),
        "saving": {
            "opex12M": numero(valor_ao_lado(ws, rot, "S. Opex 12M:")),
            "capex12M": numero(valor_ao_lado(ws, rot, "S. Capex 12M:")),
            "total12M": numero(valor_ao_lado(ws, rot, "Saving T. 12M:")),
            "opex2026": numero(valor_ao_lado(ws, rot, "S. Opex 2026:")),
            "capex2026": numero(valor_ao_lado(ws, rot, "S. Capex 2026:")),
            "total2026": numero(valor_ao_lado(ws, rot, "Saving T. 2026:")),
        },
    }
    # Macro atividades: tabela que começa no cabeçalho "MACRO ATIVIDADES:"
    macro = []
    cabecalho = None
    for r in range(1, 15):
        for c in range(1, 20):
            if chave(ws.cell(r, c).value or "") == chave("MACRO ATIVIDADES:"):
                cabecalho = (r, c)
                break
        if cabecalho:
            break
    if cabecalho:
        r0, cAtv = cabecalho
        cNum = cAtv - 1
        r = r0 + 1
        while r < r0 + 60:
            num = limpar(ws.cell(r, cNum).value)
            atividade = limpar(ws.cell(r, cAtv).value)
            if num == "" and atividade == "":
                break
            macro.append({
                "numero": texto_celula(num),
                "atividade": texto_celula(atividade),
                "responsavel": texto_celula(ws.cell(r, cAtv + 1).value),
                "status": inteiro_ou_texto(ws.cell(r, cAtv + 2).value),
                "inicio": data_iso(ws.cell(r, cAtv + 4).value),
                "fim": data_iso(ws.cell(r, cAtv + 5).value),
            })
            r += 1
    dados["macroAtividades"] = macro
    return dados


def ler_cronograma(ws):
    itens = []
    inicio = None
    for r in range(1, 15):
        if chave(ws.cell(r, 2).value or "") == chave("N°"):
            inicio = r + 1
            break
    if not inicio:
        return itens
    vazias = 0
    for r in range(inicio, ws.max_row + 1):
        num = limpar(ws.cell(r, 2).value)
        atividade = limpar(ws.cell(r, 3).value)
        if num == "" and atividade == "":
            vazias += 1
            if vazias >= 3:
                break
            continue
        vazias = 0
        num = inteiro_ou_texto(num)
        atividade = texto_celula(atividade)
        if isinstance(num, int):  # tópico principal (1, 2, 3...)
            atividade = f"{num}. {atividade}"
        itens.append({
            "numero": num if isinstance(num, int) else texto_celula(num),
            "atividade": atividade,
            "responsavel": texto_celula(ws.cell(r, 4).value),
            "status": inteiro_ou_texto(ws.cell(r, 5).value),
            "inicio": data_iso(ws.cell(r, 7).value),
            "fim": data_iso(ws.cell(r, 8).value),
        })
    return itens


def ler_anotacoes(ws):
    historico = []
    for r in range(4, ws.max_row + 1):
        data = ws.cell(r, 2).value
        acoes = limpar(ws.cell(r, 3).value)
        detalhe = limpar(ws.cell(r, 4).value)
        if data is None and detalhe == "":
            continue
        historico.append({
            "data": data_iso(data),
            "acoes": texto_celula(acoes) or "-",
            "detalhe": texto_celula(detalhe),
        })
    return historico


def ler_financeiro(ws, project_id):
    linhas = []
    for r in range(4, ws.max_row + 1):
        valores = [ws.cell(r, c).value for c in range(2, 12)]
        if not any(v not in (None, "") for v in valores):
            continue
        _, iniciativa, gerencia, projeto, situacao, icg, tipo, ano, mes, resultado = valores
        if limpar(icg) == "":
            continue
        linhas.append({
            "id": project_id,
            "iniciativa": texto_celula(iniciativa),
            "gerencia": texto_celula(gerencia),
            "projeto": texto_celula(projeto),
            "situacao": texto_celula(situacao),
            "icg": texto_celula(icg),
            "tipoDespesa": texto_celula(tipo),
            "ano": int(numero(ano)) if numero(ano) else "",
            "mes": texto_celula(mes),
            "resultado": numero(resultado),
        })
    return linhas


def aba(wb, *nomes):
    alvo = {chave(n) for n in nomes}
    for ws in wb.worksheets:
        if chave(ws.title) in alvo:
            return ws
    return None


def situacao_de(status, risco):
    """Situação usada na rosca 'Projetos por status'."""
    status = str(status).upper().strip()
    risco = chave(risco)
    if status == "L5":
        return "Concluído"
    if risco == "alto":
        return "Em risco"
    if risco == "medio":
        return "Atenção"
    if status in ("L0", "L1", "L2"):
        return "Planejado"
    return "Em execução"


def entregaveis_de(texto):
    itens = []
    for linha in str(texto or "").split("\n"):
        linha = re.sub(r"^\s*\d+\s*[.)-]\s*", "", linha).strip().rstrip(";").strip()
        if linha:
            itens.append(linha)
    return itens


def processar(caminho):
    arquivo = os.path.basename(caminho)
    project_id = normalizar_id(arquivo)
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=False)
    ws_one = aba(wb, "One_Page", "One Page")
    if ws_one is None:
        raise ValueError("aba One_Page não encontrada")
    one = ler_one_page(ws_one)
    ws_cron = aba(wb, "Cronograma")
    ws_anot = aba(wb, "Anotações", "Anotacoes")
    ws_fin = aba(wb, "Base_Financ.", "Base_Financ", "Base Financ")
    saving = one["saving"]
    tipos = [t for t, v in (("Opex", saving["opex12M"]), ("Capex", saving["capex12M"])) if v]
    projeto = {
        "id": project_id,
        "projeto": one["projeto"] or re.sub(r"^LOG-[\d.]+\s*-\s*", "", os.path.splitext(arquivo)[0]),
        "area": one["area"],
        "pontoFocal": one["pontoFocal"],
        "inicio": one["inicio"],
        "termino": one["termino"],
        "inicioCaptura": one["inicioCaptura"],
        "iniciativa": one["iniciativa"],
        "pmo": one["pmo"],
        "risco": one["risco"],
        "status": one["status"],
        "situacao": situacao_de(one["status"], one["risco"]),
        "prioridade": one["prioridade"],
        "ganhoProjeto": one["ganhoProjeto"],
        "tipos": tipos,
        "escopo": one["escopo"],
        "entregaveis": entregaveis_de(one["entregaveisTexto"]),
        "saving": saving,
        "macroAtividades": one["macroAtividades"],
        "cronogramaDetalhado": ler_cronograma(ws_cron) if ws_cron is not None else [],
        "historico": ler_anotacoes(ws_anot) if ws_anot is not None else [],
        "statusProjeto": one["statusProjeto"],
        "financeiro": ler_financeiro(ws_fin, project_id) if ws_fin is not None else [],
        "fonte": {
            "arquivo": arquivo,
            "url": urllib.parse.quote(URL_PASTA_SHAREPOINT + arquivo, safe=":/") + "?web=1",
            "extraidoEm": dt.datetime.now().astimezone().isoformat(timespec="seconds"),
        },
    }
    wb.close()
    return projeto


# ---------------------------------------------------------------------------
# Execução
# ---------------------------------------------------------------------------
def localizar_pasta_excel():
    if PASTA_EXCEL and os.path.isdir(PASTA_EXCEL):
        return PASTA_EXCEL
    raiz = os.path.expanduser("~")
    alvo = chave(NOME_PASTA_EXCEL)
    for pasta, subpastas, _ in os.walk(raiz):
        profundidade = pasta[len(raiz):].count(os.sep)
        if profundidade > 7:
            subpastas[:] = []
            continue
        subpastas[:] = [s for s in subpastas if not s.startswith(".") and s not in ("AppData", "node_modules")]
        if chave(os.path.basename(pasta)) == alvo:
            return pasta
    return ""


def main():
    pasta = localizar_pasta_excel()
    if not pasta:
        print("ERRO: pasta dos Excel não encontrada.")
        print(f'Sincronize a pasta "{NOME_PASTA_EXCEL}" no OneDrive ou preencha PASTA_EXCEL no início do gerar_jsons.py.')
        return 1
    print(f"Pasta dos projetos: {pasta}")
    os.makedirs(PASTA_JSON, exist_ok=True)

    arquivos = sorted(
        f for f in glob.glob(os.path.join(pasta, "LOG-*.xlsx"))
        if not os.path.basename(f).startswith("~$")
    )
    projetos, erros, ignorados = [], [], []
    vistos = {}
    for caminho in arquivos:
        nome = os.path.basename(caminho)
        pid = normalizar_id(nome)
        if not pid or pid.split(".")[0] in IGNORAR_IDS:
            ignorados.append(nome)
            continue
        if pid in vistos:
            erros.append(f"{nome}: ID duplicado com {vistos[pid]}")
            continue
        try:
            projeto = processar(caminho)
            projetos.append(projeto)
            vistos[pid] = nome
            print(f"  OK  {pid}  {projeto['projeto']}")
        except Exception as erro:  # segue com os demais projetos
            erros.append(f"{nome}: {erro}")
            print(f"  ERRO {nome}: {erro}")

    if not projetos:
        print("Nenhum projeto processado. Nada foi alterado na pasta json.")
        return 1

    projetos.sort(key=lambda p: ordem_id(p["id"]))
    nomes_json = [f"{p['id']}.json" for p in projetos]
    agora = dt.datetime.now().astimezone().isoformat(timespec="seconds")

    for projeto, nome in zip(projetos, nomes_json):
        with open(os.path.join(PASTA_JSON, nome), "w", encoding="utf-8") as f:
            json.dump(projeto, f, ensure_ascii=False, indent=2)

    # Remove JSONs de projetos que não existem mais (ex.: antigos LOG-15.json)
    removidos = []
    for antigo in glob.glob(os.path.join(PASTA_JSON, "LOG-*.json")):
        if os.path.basename(antigo) not in nomes_json:
            os.remove(antigo)
            removidos.append(os.path.basename(antigo))

    indice = {"schemaVersion": "1.0", "geradoEm": agora, "totalProjetos": len(nomes_json), "arquivos": nomes_json}
    with open(os.path.join(PASTA_JSON, "projects.index.json"), "w", encoding="utf-8") as f:
        json.dump(indice, f, ensure_ascii=False, indent=2)
    with open(os.path.join(PASTA_JSON, "projects.index.js"), "w", encoding="utf-8") as f:
        f.write("window.PORTFOLIO_JSON_FILES = " + json.dumps(nomes_json, ensure_ascii=False) + ";\n")
    with open(os.path.join(PASTA_JSON, "portfolio.json"), "w", encoding="utf-8") as f:
        json.dump(projetos, f, ensure_ascii=False)
    with open(os.path.join(PASTA_JSON, "projects.data.js"), "w", encoding="utf-8") as f:
        f.write("window.PORTFOLIO_DATA = " + json.dumps(projetos, ensure_ascii=False) + ";\n")

    with open(os.path.join(BASE_DIR, "MANIFESTO.txt"), "w", encoding="utf-8") as f:
        f.write("DASHBOARD PORTFÓLIO LOGÍSTICA — MANIFESTO\n\n")
        f.write(f"Projetos JSON: {len(nomes_json)}\nGerado em: {agora}\n\nProjetos:\n")
        f.writelines(f"- {n}\n" for n in nomes_json)
        if erros:
            f.write("\nArquivos com erro (não publicados):\n")
            f.writelines(f"- {e}\n" for e in erros)

    print("")
    print(f"Projetos gerados: {len(projetos)}")
    if ignorados:
        print(f"Ignorados (modelos): {', '.join(ignorados)}")
    if removidos:
        print(f"JSONs antigos removidos: {', '.join(removidos)}")
    if erros:
        print(f"ATENÇÃO - {len(erros)} arquivo(s) com erro:")
        for e in erros:
            print(f"  - {e}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
